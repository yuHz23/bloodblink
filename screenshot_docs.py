import os
import sys
import subprocess
import time

def install_deps():
    print("Installing python-docx and playwright...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "playwright"])
    print("Installing playwright browsers...")
    subprocess.check_call([sys.executable, "-m", "playwright", "install", "chromium"])

try:
    import docx
    from playwright.sync_api import sync_playwright
except ImportError:
    install_deps()
    import docx
    from playwright.sync_api import sync_playwright
from docx.shared import Inches

def main():
    doc = docx.Document()
    doc.add_heading('Hướng dẫn sử dụng BB-Console', 0)

    pages = [
        {"route": "/", "title": "1. Tổng quan (Dashboard)", "desc": "Mục đích: Giúp bạn bao quát nhanh tình hình hiện tại.\nCách dùng: Ngay khi đăng nhập, màn hình này sẽ hiện ra đầu tiên. Bạn có thể xem ngay có bao nhiêu ca cấp cứu đang cần máu, số lượng túi máu còn trong kho là bao nhiêu. Nhìn vào đây, bạn sẽ biết việc gì cần ưu tiên xử lý trước."},
        {"route": "/inventory", "title": "2. Tồn kho (Inventory)", "desc": "Mục đích: Quản lý chi tiết kho máu của bệnh viện.\nCách dùng: Màn hình này hiển thị danh sách các túi máu. Hệ thống tự động xếp các túi máu sắp hết hạn lên đầu để bạn ưu tiên sử dụng trước, tránh lãng phí. Nó cũng cảnh báo nhiệt độ bảo quản để đảm bảo an toàn cho máu."},
        {"route": "/matching", "title": "3. Ghép nguồn (Matching)", "desc": "Mục đích: Mượn hoặc xin máu từ các bệnh viện khác khi kho của bạn bị thiếu.\nCách dùng: Bạn chỉ cần chọn nhóm máu cần tìm. Hệ thống sẽ quét toàn bộ các bệnh viện trong mạng lưới và gợi ý những nơi có sẵn. Nó sẽ ưu tiên chọn bệnh viện ở gần nhất hoặc nơi có túi máu sắp hết hạn để bạn xin chuyển về."},
        {"route": "/donors", "title": "4. Tìm người hiến (Donors)", "desc": "Mục đích: Tìm danh sách những người có thể hiến máu.\nCách dùng: Tại đây, bạn sẽ thấy những ai đủ điều kiện sức khỏe và thời gian để hiến máu. Để bảo vệ sự riêng tư, hệ thống không hiện địa chỉ nhà hay vị trí chính xác của họ, mà chỉ cho biết họ đang ở khu vực nào (trong bán kính khoảng 1km)."},
        {"route": "/sos", "title": "5. SOS — Huy động", "desc": "Mục đích: Phát tín hiệu kêu gọi người hiến máu khẩn cấp.\nCách dùng: Khi phát lệnh SOS, hệ thống sẽ nhắn tin cho những người hiến máu ở gần bệnh viện nhất (1km). Nếu chưa có ai nhận lời, cứ sau mỗi 10-15 phút, hệ thống sẽ tự động gọi những người ở xa hơn (5km, 10km, 20km). Nếu quá gấp, bạn có thể tự bấm nút để gọi rộng ra ngay lập tức."},
        {"route": "/map", "title": "6. Bản đồ (Map)", "desc": "Mục đích: Xem trực quan vị trí các bệnh viện và người hiến máu.\nCách dùng: Bản đồ chấm các điểm bệnh viện (chấm to và xanh nghĩa là đang có nhiều máu). Bạn cũng sẽ thấy các vùng màu đậm nhạt thể hiện nơi nào đang có nhiều người hiến máu. Những khu vực vắng vẻ sẽ bị ẩn đi để bảo vệ thông tin cá nhân."},
        {"route": "/routing", "title": "7. Định tuyến (Routing)", "desc": "Mục đích: Dẫn đường và tính thời gian vận chuyển.\nCách dùng: Khi cần chuyển máu giữa 2 bệnh viện, tính năng này sẽ vẽ đường đi ngắn nhất và tính trước mất khoảng bao nhiêu phút để xe cứu thương chở máu tới nơi."},
        {"route": "/reports", "title": "8. Nhật ký (Chỉ dành cho Quản lý)", "desc": "Mục đích: Giám sát mọi thao tác trên hệ thống để đảm bảo tính minh bạch.\nCách dùng: Như một camera an ninh. Ghi lại chi tiết ai đã đăng nhập, ai đã xem danh tính người hiến máu. Mọi thứ được mã hóa đặc biệt để chống làm giả hoặc xóa dấu vết."}
    ]

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(viewport={"width": 1280, "height": 800})
        page = context.new_page()

        print("Navigating to login page...")
        page.goto("http://localhost:5173/")
        
        print("Logging in...")
        page.fill("#username", "admin")
        page.fill("#password", "bloodblink@2026")
        page.click("button[type='submit']")
        
        # Wait for dashboard to load
        page.wait_for_selector(".app", timeout=10000)
        time.sleep(2) # Extra wait to ensure initial data loads
        
        for item in pages:
            url = f"http://localhost:5173{item['route']}"
            print(f"Capturing {item['title']}...")
            page.goto(url)
            page.wait_for_timeout(2000) # Wait for network requests and rendering
            
            screenshot_path = f"{item['title'].split('.')[0].strip()}.png"
            page.screenshot(path=screenshot_path)
            
            doc.add_heading(item['title'], level=1)
            doc.add_paragraph(item['desc'])
            doc.add_picture(screenshot_path, width=Inches(6.0))
            
            # Clean up the screenshot file
            os.remove(screenshot_path)

        browser.close()

    output_path = r"C:\Users\Ryan\Downloads\agy\Huong_dan_su_dung_BB_Console_v2.docx"
    doc.save(output_path)
    print(f"Document saved to {output_path}")

if __name__ == "__main__":
    main()
