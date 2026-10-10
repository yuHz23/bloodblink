from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

OUT = Path(__file__).parent / 'BloodBlink_GenD_Arena_2026_VI.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY='101C2D'; NAVY2='172A40'; INK='142235'; MUTED='5B6B7D'; PAPER='F5F7FA'; WHITE='FFFFFF'
RED='D94B56'; TEAL='188F86'; BLUE='3979B8'; GOLD='D89E2A'; GREEN='28775F'
PALE_RED='FBE9EA'; PALE_TEAL='E3F5F2'; PALE_BLUE='EAF2FA'; PALE_GOLD='FFF4D8'; PALE_GREEN='E8F4ED'
LINE='DCE3EA'; FONT='Aptos'; FONT_HEAD='Aptos Display'


def rgb(value): return RGBColor.from_string(value)

def rect(slide,x,y,w,h,fill,radius=True,line_color=None):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=rgb(fill)
    if line_color: shape.line.color.rgb=rgb(line_color); shape.line.width=Pt(.8)
    else: shape.line.fill.background()
    return shape

def circle(slide,x,y,d,fill):
    shape=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d))
    shape.fill.solid(); shape.fill.fore_color.rgb=rgb(fill); shape.line.fill.background(); return shape

def line(slide,x1,y1,x2,y2,color=LINE,width=1.2):
    shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    shape.line.color.rgb=rgb(color); shape.line.width=Pt(width); return shape

def arrow(slide,x1,y1,x2,y2,color=RED,width=1.8):
    line(slide,x1,y1,x2,y2,color,width)
    if abs(x2-x1)>=abs(y2-y1): pts=[(x2-.10,y2-.07),(x2,y2),(x2-.10,y2+.07)]
    else: pts=[(x2-.07,y2-.10),(x2,y2),(x2+.07,y2-.10)]
    line(slide,*pts[0],*pts[1],color,width); line(slide,*pts[1],*pts[2],color,width)

def txt(slide,value,x,y,w,h,size=12,color=INK,bold=False,font=FONT,align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP,italic=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(0); tf.margin_right=Inches(0); tf.margin_top=Inches(0); tf.margin_bottom=Inches(0)
    tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; p.space_before=Pt(0); p.space_after=Pt(0)
    r=p.add_run(); r.text=value; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold
    r.font.italic=italic; r.font.color.rgb=rgb(color)
    return box

def rich(slide,runs,x,y,w,h,size=10,color=INK):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(0); tf.margin_right=Inches(0); tf.margin_top=Inches(0); tf.margin_bottom=Inches(0)
    p=tf.paragraphs[0]; p.space_before=Pt(0); p.space_after=Pt(0)
    for item in runs:
        r=p.add_run(); r.text=item['text']; r.font.name=FONT; r.font.size=Pt(item.get('size',size))
        r.font.bold=item.get('bold',False); r.font.color.rgb=rgb(item.get('color',color))
        if item.get('url'): r.hyperlink.address=item['url']
    return box

def card(slide,x,y,w,h,fill=WHITE,line_color=LINE): return rect(slide,x,y,w,h,fill,True,line_color)

def pill(slide,label,x,y,w,fill=PALE_BLUE,color=BLUE):
    rect(slide,x,y,w,.30,fill,True)
    txt(slide,label.upper(),x+.08,y+.035,w-.16,.22,8.5,color,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)

def bullet(slide,value,x,y,w,size=10,color=INK,dot=RED,h=.48):
    circle(slide,x,y+.085,.075,dot); txt(slide,value,x+.17,y,w-.17,h,size,color)

def base(section,title,num,subtitle=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(PAPER)
    pill(s,section,.58,.30,min(2.50,.72+len(section)*.062))
    txt(s,title,.58,.76,12.15,.55,25,INK,True,FONT_HEAD)
    if subtitle: txt(s,subtitle,.60,1.35,12,.35,10.5,MUTED); by=1.82
    else: by=1.60
    line(s,.58,7.08,12.75,7.08,LINE,.8)
    txt(s,'BLOODBLINK  •  GenD Arena 2026',.60,7.14,5.6,.18,8,MUTED)
    txt(s,f'{num:02d} / 12',11.90,7.13,.84,.18,8,MUTED,align=PP_ALIGN.RIGHT)
    return s,by

def label(s,value,x,y,w,color=BLUE): txt(s,value.upper(),x,y,w,.19,8.5,color,True)

# 1. Project overview
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(NAVY)
for x,y,d,c in [(9.7,.55,1.2,RED),(11.35,1.7,.8,TEAL),(9.35,3.1,.55,BLUE),(11.55,4.2,1.2,RED),(9.2,5.35,.42,GOLD),(12.3,.55,.34,TEAL)]: circle(s,x,y,d,c)
for a in [(10.65,1.5,11.42,1.95),(10.1,1.78,9.65,3.15),(9.88,3.55,11.65,4.38),(11.95,2.5,12.06,4.22),(9.57,3.62,9.4,5.43),(11.94,5.4,9.58,5.56)]: line(s,*a,'536A82',1.2)
pill(s,'GEN D ARENA 2026  •  Y TẾ & SỨC KHỎE',.72,.62,3.55,'263950','D4E0EC')
txt(s,'01  •  TỔNG QUAN DỰ ÁN',.77,1.29,3.4,.20,9.5,'8FD4CD',True)
txt(s,'BLOODBLINK',.72,1.62,7.7,.82,42,WHITE,True,FONT_HEAD)
txt(s,'Điều phối máu liên viện\ntừ kho đến người hiến',.75,2.62,7.65,1.30,27,'EAF0F6',True,FONT_HEAD)
txt(s,'NỀN TẢNG B2B2C  •  BLOOD BANK / HOSPITAL OPS  •  COLD-CHAIN IoT',.77,4.38,7.7,.32,10.5,'B8C7D8',True)
line(s,.77,5.12,6.85,5.12,'415168',1)
txt(s,'Một luồng điều phối: tìm kho tương hợp trước,\nchỉ huy động người hiến cho phần còn thiếu.',.78,5.40,6.35,.75,15,'D8E2ED')
txt(s,'Prototype gồm web console, backend, donor app và cảm biến chuỗi lạnh.',.78,6.55,7.0,.22,9,'91A3B8')
txt(s,'01 / 12',11.88,7.12,.88,.18,8,'91A3B8',align=PP_ALIGN.RIGHT)

# 2. Problem
s,by=base('02  •  VẤN ĐỀ','Nguồn cung lớn; điều phối vẫn là bài toán của cả mạng lưới',2,
    'Giả thuyết cần kiểm chứng: tồn kho và nhu cầu chưa đủ dễ nhìn xuyên qua các cơ sở.')
card(s,.65,by,3.05,2.32,NAVY,NAVY)
txt(s,'1,75M',.94,by+.32,2.45,.62,34,WHITE,True,FONT_HEAD)
txt(s,'đơn vị máu tiếp nhận\ntrên toàn quốc năm 2025',.96,by+1.08,2.40,.58,12,'E8EEF5')
txt(s,'94% kế hoạch năm  •  [nguồn Bộ Y tế]',.96,by+1.83,2.50,.23,8.1,'AFC0D2')
problems=[('01','TỒN KHO PHÂN TÁN','Cơ sở thiếu khó thấy sớm nguồn tương hợp đang có ở nơi khác.'),('02','ĐIỀU PHỐI RỜI RẠC','Nhóm máu, chế phẩm, hạn dùng và ETA cần được kiểm tra cùng lúc.'),('03','HUY ĐỘNG CHƯA THEO PHẦN THIẾU','SOS có nguy cơ gọi rộng nếu chưa trừ lượng kho đã điều chuyển.')]
for i,(n,h,b) in enumerate(problems):
    x=3.98+i*2.89; card(s,x,by,2.66,2.32)
    txt(s,n,x+.18,by+.18,.42,.20,9,RED,True)
    txt(s,h,x+.18,by+.52,2.29,.46,10.5,INK,True)
    txt(s,b,x+.18,by+1.08,2.29,.88,10.3,MUTED)
rect(s,.67,by+2.62,11.94,.70,PALE_GOLD,True)
txt(s,'ĐIỂM CẦN XÁC NHẬN',.89,by+2.85,1.85,.18,8.5,GOLD,True)
txt(s,'Độ trễ có đến từ phối hợp giữa các cơ sở hay không? Phỏng vấn và pilot sẽ đo trước khi kết luận.',2.92,by+2.78,9.40,.32,10.5,INK,True,valign=MSO_ANCHOR.MIDDLE)
rich(s,[{'text':'Số liệu 2025: Bộ Y tế, báo cáo Ban Chỉ đạo Quốc gia vận động hiến máu. Mở nguồn ↗','size':8,'color':BLUE,'bold':True,'url':'https://adminmoh.moh.gov.vn/web/guest/hoat-dong-cua-lanh-dao-bo/-/asset_publisher/TW6LTp1ZtwaN/content/thu-truong-bo-y-te-tran-van-thuan-ong-chu-tri-hoi-nghi-lan-thu-30-ban-chi-ao-quoc-gia-van-ong-hien-mau-tinh-nguyen?inheritRedirect=false&redirect=https%3A%2F%2Fadminmoh.moh.gov.vn%3A443%2Fweb%2Fguest%2Fhoat-dong-cua-lanh-dao-bo%3Fp_p_id%3D101_INSTANCE_TW6LTp1ZtwaN%26p_p_lifecycle%3D0%26p_p_state%3Dnormal%26p_p_mode%3Dview%26p_p_col_id%3Drow-0-column-2%26p_p_col_count%3D2'}],.73,by+3.47,11.8,.20)

# 3. Solution
s,by=base('03  •  GIẢI PHÁP','Tìm kho trước. Gọi người hiến cho phần thiếu còn lại.',3,
    'BLOODBLINK gom nhu cầu, tương hợp, offer điều chuyển, ETA và SOS vào cùng một ca.')
steps=[('01','NHẬP NHU CẦU','Cơ sở nhận • ABO/RhD • chế phẩm • số lượng • hạn cần'),('02','GHÉP KHO LIÊN VIỆN','Lọc tương hợp; xếp theo FEFO, ETA và rủi ro khan hiếm'),('03','GỬI OFFER CÓ HẠN','Cơ sở nguồn chấp nhận / từ chối; lượng nhận được giữ kho'),('04','SOS ĐÚNG PHẦN THIẾU','Offer hết hạn hoặc kho thiếu thì mở vòng gọi theo lượng còn lại')]
for i,(n,h,b) in enumerate(steps):
    x=.66+i*3.05; card(s,x,by+.28,2.76,2.50)
    circle(s,x+.20,by+.53,.48,RED if i==3 else NAVY2)
    txt(s,n,x+.20,by+.66,.48,.18,9,WHITE,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    txt(s,h,x+.20,by+1.18,2.36,.43,10.4,INK,True)
    txt(s,b,x+.20,by+1.72,2.34,.82,10.3,MUTED)
    if i<3: arrow(s,x+2.80,by+1.52,x+3.00,by+1.52,RED,1.5)
rect(s,.68,by+3.12,11.95,.62,PALE_BLUE,True)
txt(s,'NGUYÊN TẮC AN TOÀN',.90,by+3.32,1.78,.18,8.4,BLUE,True)
txt(s,'Phần mềm hỗ trợ điều phối; cơ sở y tế quyết định tương hợp cuối, xác minh và phát hành đơn vị máu.',2.78,by+3.24,9.40,.30,10.3,INK,True,valign=MSO_ANCHOR.MIDDLE)

# 4. Target customers
s,by=base('04  •  KHÁCH HÀNG MỤC TIÊU','Bắt đầu ở cụm bệnh viện có nhu cầu điều phối chung',4,
    'Khách hàng đầu tiên là tổ chức y tế; người hiến là người tham gia mạng lưới, không phải bên trả phí.')
segments=[('BÊN MUA','Mạng lưới bệnh viện\nTrung tâm / ngân hàng máu','Quản lý cần nhìn được tồn kho và năng lực điều phối nhiều cơ sở.',BLUE,PALE_BLUE),('NGƯỜI DÙNG','Điều phối truyền máu\nNhân viên kho / vận chuyển','Cần ghép tương hợp, hạn dùng, ETA và phản hồi offer trong một ca.',RED,PALE_RED),('NGƯỜI HƯỞNG LỢI','Người bệnh\nNgười hiến máu','Quy trình rõ hơn; lời mời có mục tiêu và bảo vệ quyền riêng tư.',TEAL,PALE_TEAL)]
for i,(tag,title,body,col,bg) in enumerate(segments):
    x=.68+i*4.12; card(s,x,by+.18,3.78,1.86,bg,bg)
    txt(s,tag,x+.22,by+.39,3.34,.18,8.5,col,True)
    txt(s,title,x+.22,by+.72,3.32,.50,12,INK,True)
    txt(s,body,x+.22,by+1.31,3.32,.42,9.7,MUTED)
label(s,'KHÁCH HÀNG MỤC TIÊU ĐẦU TIÊN',.79,by+2.45,3.45,BLUE)
card(s,.72,by+2.82,7.18,1.05)
txt(s,'Cụm bệnh viện đô thị có khoa truyền máu / kho máu và đầu mối điều phối.',.95,by+3.03,6.72,.55,13,INK,True,valign=MSO_ANCHOR.MIDDLE)
card(s,8.14,by+2.82,4.45,1.05,PALE_GOLD,PALE_GOLD)
txt(s,'Giả thuyết pilot: 3 cơ sở neo tại TP.HCM → mở rộng tối đa 8 cơ sở.',8.40,by+3.02,3.90,.56,10.5,GOLD,True,valign=MSO_ANCHOR.MIDDLE)

# 5. MVP / Prototype
s,by=base('05  •  MVP / PROTOTYPE','MVP có đủ bốn mô-đun để diễn tập một luồng điều phối',5,
    'MVP là prototype trong repository — chưa phải hệ thống đã được bệnh viện xác nhận trong vận hành thật.')
mods=[('BB-CONSOLE','Bản đồ kho, ca SOS, đề nghị điều chuyển, phản hồi offer và tuyến ETA.',BLUE,PALE_BLUE),('BB-CORE','API, ABO/RhD, FEFO, offer có hạn, lịch SOS, phân quyền và audit.',RED,PALE_RED),('BB-SOS','Hồ sơ người hiến, hộp thư SOS, nhận / từ chối lời mời, cơ sở gần nhất.',TEAL,PALE_TEAL),('BB-BOX','Firmware ESP32-S3 đọc cảm biến; ngưỡng bảo quản và vận chuyển tách riêng.',GOLD,PALE_GOLD)]
for i,(name,body,col,bg) in enumerate(mods):
    x=.70+(i%2)*6.10; y=by+.10+(i//2)*1.18; card(s,x,y,5.80,.98,WHITE)
    rect(s,x,y,.08,.98,col,False); txt(s,name,x+.22,y+.18,1.42,.20,9,col,True)
    txt(s,body,x+1.66,y+.16,3.88,.62,9.7,INK)
rect(s,.72,by+2.64,11.91,1.06,NAVY,True)
txt(s,'ĐÃ CÓ TRONG CODE',.97,by+2.86,2.15,.18,8.5,'B8C8D8',True)
txt(s,'Web map • backend • mobile flows • sensor firmware • simulation',.97,by+3.17,5.36,.26,10.5,WHITE,True)
txt(s,'CÒN PHẢI KIỂM CHỨNG',6.82,by+2.86,2.50,.18,8.5,'FFC8CA',True)
txt(s,'HIS/LIS integration • FCM • sensor calibration • clinical workflow • real pilot',6.82,by+3.17,5.32,.30,9.9,WHITE,True)
txt(s,'README repo báo cáo 310 test pass; đây không phải xác nhận lâm sàng hay nghiệm thu triển khai.',.84,by+3.93,11.7,.22,9,MUTED,italic=True,align=PP_ALIGN.CENTER)

# 6. Technology requirements
s,by=base('06  •  CÔNG NGHỆ','Công nghệ được gắn với từng bước — không dùng AI để quyết định',6,
    'Trọng tâm: ghép đúng, điều phối có thời hạn, giữ riêng tư và giám sát chuỗi lạnh.')
# table header and rows
headers=[('CÔNG NGHỆ',.75,2.15),('ÁP DỤNG Ở ĐÂU',3.04,2.55),('GIẢI QUYẾT / TẠO KHÁC BIỆT GÌ',5.75,3.45),('ĐỘ TRƯỞNG THÀNH',9.38,3.02)]
for h,x,w in headers: rect(s,x,by+.06,w,.40,NAVY,True); txt(s,h,x+.08,by+.15,w-.16,.18,8,WHITE,True,align=PP_ALIGN.CENTER)
rows=[
 ('React + TypeScript','BB-Console web map','Một màn hình cho kho, ca, offer, ETA.','Prototype UI',BLUE),
 ('FastAPI + PostgreSQL/PostGIS','BB-Core API','Dữ liệu nhiều cơ sở, phân quyền và workflow có audit.','Prototype backend',RED),
 ('Rule engine ABO/RhD + FEFO + ETA','Trước khi gửi offer','Xếp hạng tương hợp, hạn dùng, ETA và rủi ro O−.','Logic xác định; cần chuyên gia duyệt',TEAL),
 ('Flutter + H3 + ESP32-S3 / TMP117','BB-SOS + BB-Box','Ứng dụng người hiến; vị trí dạng ô; hai chế độ ngưỡng nhiệt.','Prototype; tích hợp/hiệu chuẩn còn mở',GOLD)]
for i,(tech,where,why,maturity,col) in enumerate(rows):
    y=by+.57+i*.69
    if i%2==0: rect(s,.74,y,11.68,.59,WHITE,True)
    txt(s,tech,.87,y+.11,2.04,.39,8.9,col,True)
    txt(s,where,3.14,y+.11,2.34,.39,8.9,INK,True)
    txt(s,why,5.86,y+.10,3.28,.43,8.8,INK)
    txt(s,maturity,9.48,y+.10,2.72,.43,8.4,MUTED)
rect(s,.76,by+3.49,11.65,.58,PALE_BLUE,True)
txt(s,'AI',.98,by+3.68,.35,.18,9,BLUE,True)
txt(s,'Không dùng AI/ML trong ghép nguồn hay xác định đủ điều kiện; quyết định dùng quy tắc minh bạch ở backend.',1.48,by+3.61,10.60,.28,9.5,INK,True,valign=MSO_ANCHOR.MIDDLE)
rich(s,[{'text':'Thông tư 26/2013/TT-BYT — nguồn đối chiếu quy trình truyền máu. Mở CSDL VBPL ↗','size':8,'color':BLUE,'bold':True,'url':'https://vbpl.vn/boyte/Pages/vbpq-toanvan.aspx?ItemID=46980'}],.84,by+4.18,11.5,.18)

# 7. Competitors / advantage
s,by=base('07  •  ĐỐI THỦ & LỢI THẾ CẠNH TRANH','Lợi thế là khép vòng phối hợp, không phải một thuật toán đơn lẻ',7,
    'So sánh theo nhóm giải pháp thay thế; cần phỏng vấn người mua và khảo sát vendor trước khi kết luận thị trường.')
# matrix
cols=[.75,3.13,5.48,7.81,10.20]; widths=[2.18,2.15,2.13,2.18,2.23]
heads=['BƯỚC QUY TRÌNH','GỌI / NHÓM CHAT','HỆ THỐNG TẠI CHỖ','APP NGƯỜI HIẾN','BLOODBLINK']
for j,(h,x,w) in enumerate(zip(heads,cols,widths)):
    rect(s,x,by+.10,w,.57,NAVY if j==4 else 'E6ECF2',True)
    txt(s,h,x+.06,by+.22,w-.12,.31,8.2,WHITE if j==4 else MUTED,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
rows=[('Nhìn tồn kho xuyên cơ sở','Thủ công','Thường cục bộ*','—','Có luồng liên viện'),('Ghép ABO/RhD + FEFO + ETA','Rời rạc','Có thể nội bộ*','—','Một bộ xếp hạng'),('Offer có thời hạn + giữ kho','Gọi riêng','—','—','Có workflow'),('SOS sau khi trừ lượng kho','Có thể gọi rộng','—','Có thể gọi','Chỉ phần thiếu'),('Tín hiệu chuỗi lạnh','Tách biệt','Có thể có*','—','BB-Box module')]
for i,row in enumerate(rows):
    y=by+.76+i*.53
    if i%2==0: rect(s,.75,y,11.69,.46,WHITE,True)
    txt(s,row[0],.88,y+.12,2.03,.22,8.6,INK,True)
    for j in range(1,5): txt(s,row[j],cols[j]+.06,y+.12,widths[j]-.12,.22,8.3,TEAL if j==4 else MUTED,j==4,align=PP_ALIGN.CENTER)
rect(s,.76,by+3.56,11.68,.48,PALE_GOLD,True)
txt(s,'ĐIỂM KHÁC BIỆT',.94,by+3.71,1.65,.17,8.2,GOLD,True)
txt(s,'Kho trước → offer liên viện → SOS phần thiếu → audit / cold-chain trong cùng ca.',2.72,by+3.65,9.35,.25,9.5,INK,True,valign=MSO_ANCHOR.MIDDLE)
txt(s,'*Phạm vi thay đổi theo cơ sở; chưa phải so sánh sản phẩm cụ thể.',.88,by+4.13,11.2,.18,8.4,MUTED,italic=True)

# 8. Business Model Canvas
s,by=base('08  •  BUSINESS MODEL CANVAS','BMC: giá trị cho mạng lưới y tế, doanh thu từ tổ chức',8,
    'Bản mô hình ban đầu để kiểm chứng qua pilot; chưa có khách hàng trả phí hoặc hợp đồng trong repository.')
bmc=[
 ('ĐỐI TÁC CHÍNH','Bệnh viện / trung tâm máu\nCố vấn truyền máu\nĐối tác thiết bị & tích hợp',BLUE),
 ('HOẠT ĐỘNG CHÍNH','Ghép kho, ETA, offer\nVận hành SOS\nTích hợp, bảo mật, hỗ trợ',RED),
 ('GIÁ TRỊ CỐT LÕI','Tìm kho trước, gọi sau\nMột luồng có truy vết\nẨn danh vị trí người hiến',TEAL),
 ('QUAN HỆ KHÁCH HÀNG','Onboarding theo cơ sở\nĐào tạo điều phối\nHỗ trợ & SLA theo hợp đồng',GOLD),
 ('PHÂN KHÚC KHÁCH HÀNG','Mạng lưới bệnh viện\nTrung tâm / ngân hàng máu\nBên mua: tổ chức y tế',BLUE),
 ('TÀI NGUYÊN CHÍNH','BB-Core + BB-Console\nBB-SOS + BB-Box\nQuy tắc & dữ liệu vận hành',RED),
 ('KÊNH','Bán hàng trực tiếp B2B\nDemo tại bệnh viện\nHội nghị / hội chuyên môn',TEAL),
 ('DOANH THU','Thuê bao theo cơ sở/mạng lưới\nPhí triển khai & tích hợp\nBB-Box / bảo trì tùy chọn',GOLD),
 ('CẤU TRÚC CHI PHÍ','Kỹ thuật & bảo mật\nTích hợp / hỗ trợ pilot\nThiết bị, cố vấn, bán hàng',BLUE)]
for i,(h,b,col) in enumerate(bmc):
    c=i%3; r=i//3; x=.70+c*4.10; y=by+.03+r*1.24
    card(s,x,y,3.80,1.06,WHITE)
    rect(s,x,y,.07,1.06,col,False)
    txt(s,h,x+.20,y+.14,3.40,.17,8.2,col,True)
    txt(s,b,x+.20,y+.41,3.38,.53,8.7,INK)
rect(s,.72,by+3.80,11.86,.43,PALE_GOLD,True)
txt(s,'GIẢ ĐỊNH GIÁ & QUY MÔ CẦN KIỂM CHỨNG; không tính phí người hiến và không bán dữ liệu cá nhân.',.91,by+3.93,11.44,.18,8.6,GOLD,True,align=PP_ALIGN.CENTER)

# 9. Marketing and communications
s,by=base('09  •  MARKETING & TRUYỀN THÔNG','Bán bằng độ tin cậy chuyên môn và bằng chứng pilot',9,
    'Kế hoạch truyền thông B2B y tế; không quảng cáo kết quả lâm sàng chưa được đo.')
# audience/messaging
card(s,.72,by+.08,4.0,1.13,NAVY,NAVY)
txt(s,'THÔNG ĐIỆP CHÍNH',.96,by+.28,3.48,.19,8.5,'B8C8D8',True)
txt(s,'“Kho trước. Gọi đúng phần thiếu.”',.96,by+.62,3.45,.43,15,WHITE,True,FONT_HEAD)
channels=[('B2B TRỰC TIẾP','Lãnh đạo bệnh viện, trung tâm máu, khoa truyền máu.'),('DEMO / HỘI CHUYÊN MÔN','Demo luồng offer và SOS; hội thảo chuyển đổi số y tế.'),('NỘI DUNG CÓ BẰNG CHỨNG','Case study sau pilot; chỉ công bố KPI đã được cơ sở duyệt.')]
for i,(h,b) in enumerate(channels):
    y=by+.08+i*.83
    card(s,5.00,y,7.60,.67,WHITE)
    txt(s,h,5.22,y+.13,2.18,.18,8.5,BLUE if i==0 else (TEAL if i==1 else RED),True)
    txt(s,b,7.48,y+.10,4.85,.38,9,INK)
# campaign timeline
label(s,'NHỊP 12 THÁNG',.79,by+2.83,1.8,BLUE)
phases=[('M1–3','15 phỏng vấn mục tiêu\n3 bệnh viện neo'),('M4–6','Demo kín + pilot 3 cơ sở\nĐào tạo người dùng'),('M7–12','Báo cáo pilot đã duyệt\nMở rộng mục tiêu 8 cơ sở')]
for i,(t,b) in enumerate(phases):
    x=.75+i*4.10; card(s,x,by+3.18,3.72,.91)
    pill(s,t,x+.14,by+3.31,.80,PALE_BLUE,BLUE)
    txt(s,b,x+1.10,by+3.31,2.43,.54,9.1,INK,True)
text='KPI funnel: tổ chức đủ điều kiện → cuộc gặp → demo → pilot → gia hạn. Mọi kết quả mô phỏng phải được phân biệt với số liệu thực.'
txt(s,text,.83,by+4.22,11.65,.28,8.8,MUTED,italic=True,align=PP_ALIGN.CENTER)

# 10. Financial plan
s,by=base('10  •  KẾ HOẠCH TÀI CHÍNH','Kịch bản 12 tháng: chi phí pilot lớn hơn doanh thu ban đầu',10,
    'Giả định đề xuất — chưa có báo giá, hợp đồng hay doanh thu thực tế trong repository.')
card(s,.68,by+.10,5.60,2.60,WHITE)
label(s,'DOANH THU GIẢ ĐỊNH',.96,by+.32,2.75,BLUE)
txt(s,'Thuê bao: 8 cơ sở × 6 tháng × 5M',.97,by+.75,4.10,.23,9.5,INK)
txt(s,'240M',5.08,by+.71,.87,.30,14,BLUE,True,align=PP_ALIGN.RIGHT)
txt(s,'Triển khai: 8 × 15M / cơ sở',.97,by+1.14,4.10,.23,9.5,INK)
txt(s,'120M',5.08,by+1.10,.87,.30,14,BLUE,True,align=PP_ALIGN.RIGHT)
line(s,.97,by+1.49,5.95,by+1.49,LINE,.8)
txt(s,'TỔNG DOANH THU',.97,by+1.71,3.18,.24,9.2,BLUE,True)
txt(s,'360M',4.82,by+1.66,1.12,.32,16,BLUE,True,align=PP_ALIGN.RIGHT)
rect(s,.95,by+2.13,5.00,.35,PALE_BLUE,True)
txt(s,'MRR tại 8 cơ sở: 40M  •  ARR run-rate: 480M',1.08,by+2.22,4.73,.16,8.7,BLUE,True,align=PP_ALIGN.CENTER)
card(s,6.48,by+.10,6.14,2.60,WHITE)
label(s,'CHI PHÍ 12 THÁNG',6.76,by+.32,2.55,RED)
costs=[('Sản phẩm & kỹ thuật','480M'),('Tích hợp / triển khai','120M'),('Cloud, bảo mật, vận hành','96M'),('8 thiết bị pilot','32M'),('Marketing, đào tạo, đi lại','96M'),('Cố vấn lâm sàng / pháp lý','48M'),('Dự phòng 10%','87M')]
for i,(a,b) in enumerate(costs):
    y=by+.72+i*.245; txt(s,a,6.77,y,4.12,.18,8.8,INK); txt(s,b,11.50,y,.78,.18,8.8,INK,True,align=PP_ALIGN.RIGHT)
rect(s,.70,by+2.96,11.90,.68,NAVY,True)
txt(s,'LỖ DỰ KIẾN',.92,by+3.18,1.45,.18,8.5,'B8C8D8',True)
txt(s,'~599M',2.47,by+3.10,1.40,.31,15,WHITE,True)
txt(s,'ĐỀ NGHỊ VỐN / GRANT',4.53,by+3.18,2.00,.18,8.5,'FFC8CA',True)
txt(s,'800M',6.65,by+3.10,1.28,.31,15,'FFC8CA',True)
txt(s,'~600M gap + ~200M vốn lưu động',8.29,by+3.15,3.80,.22,9,WHITE,True)
txt(s,'Năm pilot chưa hòa vốn. Giá, chi phí và chu kỳ mua sắm phải xác nhận bằng báo giá / cam kết thực.',.84,by+3.90,11.55,.23,8.8,MUTED,italic=True,align=PP_ALIGN.CENTER)

# 11. Roadmap + pivot
s,by=base('11  •  ROADMAP & PIVOT','Mở rộng theo bằng chứng; pivot theo tín hiệu người mua',11,
    'Các mốc là kế hoạch đề xuất, không phải lịch sử triển khai đã hoàn thành.')
phases=[('0–3 THÁNG','CHUẨN BỊ','3 bệnh viện neo\nCố vấn truyền máu\nRà soát dữ liệu / bảo mật',BLUE),('4–6 THÁNG','MVP PILOT','Chạy 3 cơ sở\nĐo baseline + KPI\nHiệu chuẩn BB-Box',TEAL),('7–12 THÁNG','MỞ RỘNG CÓ ĐIỀU KIỆN','Tối đa 8 cơ sở\nĐánh giá gia hạn\nBáo cáo tác động',RED)]
for i,(period,h,b,col) in enumerate(phases):
    x=.72+i*4.08; card(s,x,by+.08,3.70,1.36,WHITE)
    pill(s,period,x+.17,by+.25,1.38,PALE_BLUE,BLUE)
    txt(s,h,x+.17,by+.69,3.35,.20,9.5,INK,True)
    txt(s,b,x+.17,by+.98,3.35,.38,8.8,MUTED)
label(s,'PIVOT GATES — ĐỀ XUẤT',.80,by+1.78,2.65,GOLD)
triggers=[('DỮ LIỆU / INTEGRATION CHẬM','Nếu ≤1/3 cơ sở pilot chia sẻ được dữ liệu tối thiểu: thu hẹp MVP về inventory + offer; hoãn SOS/IoT.'),('NGƯỜI MUA TẬP TRUNG','Nếu một hệ thống quyết định cho nhiều cơ sở: pivot từ thuê bao từng cơ sở sang license mạng lưới.'),('KHÔNG CÓ WILLINGNESS-TO-PAY','Nếu pilot không có gia hạn trả phí: kiểm tra buyer/funding model trước khi mở rộng đội ngũ.')]
for i,(h,b) in enumerate(triggers):
    y=by+2.12+i*.67; card(s,.76,y,11.79,.54,PALE_GOLD,PALE_GOLD)
    txt(s,h,.94,y+.10,2.86,.18,8.1,GOLD,True)
    txt(s,b,3.95,y+.08,8.30,.36,8.8,INK)
txt(s,'Pivot là thay đổi phạm vi / kênh / mô hình doanh thu dựa trên dữ liệu pilot — không đổi quy tắc an toàn truyền máu.',.86,by+4.20,11.53,.25,8.8,MUTED,italic=True,align=PP_ALIGN.CENTER)

# 12. Team and conclusion
s,by=base('12  •  TEAM / KẾT LUẬN','Đã có prototype kỹ thuật; cần đúng đối tác để pilot an toàn',12,
    'Không nêu tên cá nhân; tập trung vào năng lực triển khai cần có cho giai đoạn tiếp theo.')
card(s,.72,by+.12,5.80,1.22,PALE_BLUE,PALE_BLUE)
txt(s,'NĂNG LỰC ĐÃ THỂ HIỆN QUA SẢN PHẨM',.98,by+.34,5.20,.18,8.5,BLUE,True)
txt(s,'Web + backend • donor app • embedded IoT • simulation',.98,by+.72,5.15,.34,11,INK,True)
card(s,6.80,by+.12,5.80,1.22,PALE_TEAL,PALE_TEAL)
txt(s,'NĂNG LỰC CẦN ĐỒNG HÀNH',7.06,by+.34,5.18,.18,8.5,TEAL,True)
txt(s,'Chuyên gia truyền máu • vận hành bệnh viện • bảo mật dữ liệu',7.06,by+.72,5.15,.34,10.5,INK,True)
# close / ask
rect(s,.72,by+1.78,11.88,1.78,NAVY,True)
txt(s,'ĐỀ NGHỊ',.99,by+2.04,1.05,.20,8.5,'B8C8D8',True)
txt(s,'3 bệnh viện pilot  +  cố vấn chuyên môn  +  800M vốn / grant cho 12 tháng',.99,by+2.39,11.10,.42,17,WHITE,True,FONT_HEAD)
txt(s,'Kho trước. Gọi đúng phần thiếu. Đo trước khi mở rộng.',.99,by+3.05,11.10,.28,11,'FFC8CA',True)
txt(s,'BLOODBLINK  |  GenD Arena 2026  |  Nhiệm vụ 1: Pitch Deck',.83,by+3.95,11.55,.22,9,MUTED,align=PP_ALIGN.CENTER)

prs.core_properties.title='BLOODBLINK — GenD Arena 2026 Pitch Deck (12 nội dung bắt buộc)'
prs.core_properties.subject='Nhiệm vụ 1 — Pitch Deck'
prs.core_properties.author='BLOODBLINK project'
prs.core_properties.keywords='BloodBlink, MVP, blood coordination, GenD Arena 2026'
prs.save(OUT)
print(f'Created {OUT} ({len(prs.slides)} slides)')
